import os
import json
import re
import math
import secrets
import hashlib
import time
from datetime import datetime, timezone
import tempfile
import asyncio
import uuid
from io import BytesIO
from fastapi import FastAPI, Depends, HTTPException, File, UploadFile, Form, WebSocket, WebSocketDisconnect, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from google import genai
from google.genai import types
from groq import Groq
from dotenv import load_dotenv

import models, schemas
from database import engine, get_db
from cache import (
    ai_cache,
    AI_CONCURRENCY_SEMAPHORE,
    make_deck_cache_key,
    make_chat_cache_key,
    make_readout_cache_key,
)
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

load_dotenv()

# --- INITIALIZE BOTH AI ENGINES ---
gemini_api_key = os.environ.get("GEMINI_API_KEY") or os.environ.get("GOOGLE_API_KEY")
try:
    if gemini_api_key:
        ai_client = genai.Client(api_key=gemini_api_key)
    else:
        ai_client = genai.Client()
except Exception as e:
    print(f"Warning: Gemini Client failed to initialize. {e}")
    ai_client = None

try:
    groq_api_key = os.environ.get("GROQ_API_KEY")
    groq_client = Groq(api_key=groq_api_key) if groq_api_key else None
except Exception as e:
    print(f"Warning: Groq Client failed to initialize. {e}")
    groq_client = None

GROQ_MODEL = os.environ.get("GROQ_MODEL", "openai/gpt-oss-120b")
FALLBACK_GEMINI_MODELS = ["gemini-2.0-flash", "gemini-1.5-flash"]


def extract_json_data(text: str):
    if not text:
        raise ValueError("Empty response from AI")
    cleaned = text.strip()
    match = re.search(r"(\[[\s\S]*\]|\{[\s\S]*\})", cleaned)
    if match:
        return json.loads(match.group(1))
    return json.loads(cleaned)

from sqlalchemy import text
from sqlalchemy.orm import Session

models.Base.metadata.create_all(bind=engine)

# Auto-migrate: ensure user_id and author_name columns exist on decks table in Neon (PostgreSQL) / SQLite
try:
    with engine.connect() as conn:
        if engine.dialect.name == "sqlite":
            res = conn.execute(text("PRAGMA table_info(decks)")).fetchall()
            cols = [row[1] for row in res]
            if "user_id" not in cols:
                conn.execute(text("ALTER TABLE decks ADD COLUMN user_id VARCHAR(64)"))
                conn.commit()
            if "author_name" not in cols:
                conn.execute(text("ALTER TABLE decks ADD COLUMN author_name VARCHAR(120) DEFAULT 'Student'"))
                conn.commit()
        else:
            conn.execute(text("ALTER TABLE decks ADD COLUMN IF NOT EXISTS user_id VARCHAR(64)"))
            conn.execute(text("ALTER TABLE decks ADD COLUMN IF NOT EXISTS author_name VARCHAR(120) DEFAULT 'Student'"))
            conn.commit()
except Exception as e:
    print(f"Migration note (columns on decks): {e}")

app = FastAPI(title="CramZero API")

# Configure CORS origins - allow Vercel deployments, custom domains, and local dev
allowed_origins_env = os.getenv("ALLOWED_ORIGINS") or os.getenv("FRONTEND_URL")
allowed_origins = [
    "https://cram-zero.vercel.app",
    "http://localhost:5173",
    "http://localhost:3000",
    "http://localhost:4173",
    "http://127.0.0.1:5173",
    "http://127.0.0.1:3000",
]
if allowed_origins_env:
    for origin in allowed_origins_env.split(","):
        cleaned = origin.strip().rstrip("/")
        if cleaned and cleaned not in allowed_origins:
            allowed_origins.append(cleaned)

app.add_middleware(
    CORSMiddleware,
    allow_origins=allowed_origins,
    allow_origin_regex=r"^https:\/\/.*\.vercel\.app$",
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.get("/api/health")
async def health_check():
    return {"status": "CramZero Hybrid Engines online"}

# --- RATE LIMITING ENGINE ---
# Configurable daily limits (reset daily at 00:00 UTC)
RATE_LIMITS = {
    "deck_generation": {
        "limit": 5,
        "label": "deck generations",
    },
    "chat": {
        "limit": 30,
        "label": "Zero AI chats",
    },
    "readout": {
        "limit": 10,
        "label": "reading guide PDF exports",
    },
}

_in_memory_usage = {}

def get_current_date_str() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d")

def get_client_identifier(request: Request, user_id: str | None = None) -> str:
    if user_id and str(user_id).strip():
        return f"user:{str(user_id).strip()}"
    if request:
        forwarded = request.headers.get("x-forwarded-for")
        if forwarded:
            client_ip = forwarded.split(",")[0].strip()
            return f"ip:{client_ip}"
        if request.client and request.client.host:
            return f"ip:{request.client.host}"
    return "ip:anonymous"

def get_usage_count(action: str, identifier: str, db: Session | None = None) -> int:
    date_str = get_current_date_str()
    key = (identifier, action, date_str)
    
    if db:
        try:
            record = db.query(models.ApiUsage).filter(
                models.ApiUsage.identifier == identifier,
                models.ApiUsage.action == action,
                models.ApiUsage.date_str == date_str,
            ).first()
            if record:
                _in_memory_usage[key] = record.count
                return record.count
        except Exception:
            pass
            
    return _in_memory_usage.get(key, 0)

def enforce_rate_limit(action: str, identifier: str, db: Session | None = None):
    config = RATE_LIMITS.get(action)
    if not config:
        return
        
    limit = config["limit"]
    label = config["label"]
    current_count = get_usage_count(action, identifier, db)
    
    if current_count >= limit:
        raise HTTPException(
            status_code=429,
            detail=f"Daily limit reached for {label} ({current_count}/{limit}). Limit resets tomorrow at 00:00 UTC to protect server resources."
        )

def record_rate_limit_usage(action: str, identifier: str, db: Session | None = None):
    config = RATE_LIMITS.get(action)
    if not config:
        return
        
    date_str = get_current_date_str()
    key = (identifier, action, date_str)
    
    new_count = _in_memory_usage.get(key, 0) + 1
    _in_memory_usage[key] = new_count
    
    if db:
        try:
            record = db.query(models.ApiUsage).filter(
                models.ApiUsage.identifier == identifier,
                models.ApiUsage.action == action,
                models.ApiUsage.date_str == date_str,
            ).first()
            if record:
                record.count += 1
            else:
                record = models.ApiUsage(
                    identifier=identifier,
                    action=action,
                    date_str=date_str,
                    count=1,
                )
                db.add(record)
            db.commit()
            _in_memory_usage[key] = record.count
        except Exception as e:
            db.rollback()
            print(f"Rate limit DB record notice: {e}")

@app.get("/api/rate-limits/")
def get_rate_limits_status(req: Request, user_id: str | None = None, db: Session = Depends(get_db)):
    identifier = get_client_identifier(req, user_id)
    status = {}
    for action, config in RATE_LIMITS.items():
        used = get_usage_count(action, identifier, db)
        limit = config["limit"]
        status[action] = {
            "used": used,
            "limit": limit,
            "remaining": max(0, limit - used),
            "label": config["label"],
        }
    return status

# --- PASSWORD HASHING UTILITIES ---
def hash_password(password: str) -> str:
    salt = secrets.token_hex(16)
    key = hashlib.pbkdf2_hmac("sha256", password.encode(), salt.encode(), 100000)
    return f"{salt}${key.hex()}"

def verify_password(password: str, hashed: str | None) -> bool:
    if not hashed or "$" not in hashed:
        return False
    try:
        salt, key = hashed.split("$", 1)
        check = hashlib.pbkdf2_hmac("sha256", password.encode(), salt.encode(), 100000)
        return secrets.compare_digest(check.hex(), key)
    except Exception:
        return False

def _user_payload(user):
    return {
        "id": user.id,
        "email": user.email,
        "display_name": user.display_name or "Student",
        "avatar_url": user.avatar_url,
        "is_guest": not bool(user.email and user.password_hash),
        "created_at": user.created_at,
    }

# --- AUTH & USER PROFILE ENDPOINTS ---
@app.post("/api/auth/register", response_model=schemas.UserResponse)
def register_user(payload: schemas.UserRegister, db: Session = Depends(get_db)):
    email = payload.email.strip().lower()
    if not email or "@" not in email:
        raise HTTPException(status_code=400, detail="Please enter a valid email address.")
    if not payload.password or len(payload.password) < 6:
        raise HTTPException(status_code=400, detail="Password must be at least 6 characters.")
    
    existing = db.query(models.User).filter(models.User.email == email).first()
    if existing:
        raise HTTPException(status_code=400, detail="An account with this email already exists. Please sign in.")
    
    target_user_id = payload.guest_user_id.strip() if payload.guest_user_id else None
    user = None
    if target_user_id:
        user = db.query(models.User).filter(models.User.id == target_user_id).first()
    
    display_name = payload.display_name.strip() if payload.display_name else "Student"

    if user:
        user.email = email
        user.password_hash = hash_password(payload.password)
        user.display_name = display_name
        if payload.avatar_url:
            user.avatar_url = payload.avatar_url
    else:
        user_id = target_user_id or f"user-{secrets.token_hex(12)}"
        user = models.User(
            id=user_id,
            email=email,
            password_hash=hash_password(payload.password),
            display_name=display_name,
            avatar_url=payload.avatar_url,
        )
        db.add(user)
    
    # Ensure all existing guest decks & hub memberships are tied to this user
    if target_user_id and target_user_id != user.id:
        db.query(models.Deck).filter(models.Deck.user_id == target_user_id).update({"user_id": user.id, "author_name": user.display_name})
        db.query(models.StudyHubMember).filter(models.StudyHubMember.user_id == target_user_id).update({"user_id": user.id, "display_name": user.display_name})
    
    db.commit()
    db.refresh(user)
    return _user_payload(user)


@app.post("/api/auth/login", response_model=schemas.UserResponse)
def login_user(payload: schemas.UserLogin, db: Session = Depends(get_db)):
    email = payload.email.strip().lower()
    user = db.query(models.User).filter(models.User.email == email).first()
    if not user or not verify_password(payload.password, user.password_hash):
        raise HTTPException(status_code=401, detail="Invalid email or password.")
    
    # If the user created decks as guest on this device, merge them into their logged in account!
    if payload.guest_user_id and payload.guest_user_id != user.id:
        db.query(models.Deck).filter(models.Deck.user_id == payload.guest_user_id).update({"user_id": user.id, "author_name": user.display_name})
        db.query(models.StudyHubMember).filter(models.StudyHubMember.user_id == payload.guest_user_id).update({"user_id": user.id, "display_name": user.display_name})
        db.commit()

    return _user_payload(user)


@app.post("/api/auth/guest", response_model=schemas.UserResponse)
def setup_guest_user(payload: schemas.GuestUserCreate, db: Session = Depends(get_db)):
    user_id = (payload.user_id.strip() if payload.user_id else None) or f"guest-{secrets.token_hex(12)}"
    display_name = payload.display_name.strip() if payload.display_name else "Student"
    
    user = db.query(models.User).filter(models.User.id == user_id).first()
    if user:
        user.display_name = display_name
        if payload.avatar_url:
            user.avatar_url = payload.avatar_url
    else:
        user = models.User(
            id=user_id,
            display_name=display_name,
            avatar_url=payload.avatar_url,
        )
        db.add(user)
    
    db.query(models.StudyHubMember).filter(models.StudyHubMember.user_id == user.id).update({"display_name": display_name})
    db.commit()
    db.refresh(user)
    return _user_payload(user)


@app.get("/api/auth/me", response_model=schemas.UserResponse)
def get_current_user_profile(user_id: str, db: Session = Depends(get_db)):
    user = db.query(models.User).filter(models.User.id == user_id).first()
    if not user:
        user = models.User(id=user_id, display_name="Student")
        db.add(user)
        db.commit()
        db.refresh(user)
    return _user_payload(user)


@app.put("/api/auth/profile", response_model=schemas.UserResponse)
def update_user_profile(payload: schemas.UserProfileUpdate, db: Session = Depends(get_db)):
    user = db.query(models.User).filter(models.User.id == payload.user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found.")
    
    if payload.display_name:
        user.display_name = payload.display_name.strip()
        db.query(models.StudyHubMember).filter(models.StudyHubMember.user_id == user.id).update({"display_name": user.display_name})
        db.query(models.Deck).filter(models.Deck.user_id == user.id).update({"author_name": user.display_name})
    
    if payload.avatar_url is not None:
        user.avatar_url = payload.avatar_url
    
    # Check if upgrading guest to full account
    if payload.email and payload.password:
        email = payload.email.strip().lower()
        if not email or "@" not in email:
            raise HTTPException(status_code=400, detail="Invalid email format.")
        if len(payload.password) < 6:
            raise HTTPException(status_code=400, detail="Password must be at least 6 characters.")
        existing = db.query(models.User).filter(models.User.email == email, models.User.id != user.id).first()
        if existing:
            raise HTTPException(status_code=400, detail="An account with this email already exists.")
        user.email = email
        user.password_hash = hash_password(payload.password)
    
    db.commit()
    db.refresh(user)
    return _user_payload(user)


# --- DECK MANAGEMENT & SHARING ENDPOINTS ---
@app.get("/api/decks/", response_model=list[schemas.DeckResponse])
def get_decks(user_id: str | None = None, db: Session = Depends(get_db)):
    if user_id:
        return db.query(models.Deck).filter(models.Deck.user_id == user_id).order_by(models.Deck.created_at.desc()).all()
    return []


@app.get("/api/decks/public")
def get_public_decks(search: str | None = None, limit: int = 50, db: Session = Depends(get_db)):
    query = db.query(models.Deck).filter(models.Deck.visibility == "Public")
    if search and search.strip():
        term = f"%{search.strip()}%"
        query = query.filter(models.Deck.title.ilike(term) | models.Deck.description.ilike(term))
    decks = query.order_by(models.Deck.created_at.desc()).limit(limit).all()
    return [_deck_payload(d) for d in decks]


@app.get("/api/decks/{deck_id}")
def get_single_deck(deck_id: int, user_id: str | None = None, db: Session = Depends(get_db)):
    deck = db.query(models.Deck).filter(models.Deck.id == deck_id).first()
    if not deck:
        raise HTTPException(status_code=404, detail="Deck not found.")
    if deck.visibility != "Public" and deck.user_id != user_id:
        # Check if in a shared study hub with user
        shared = db.query(models.StudyHubDeck).join(
            models.StudyHubMember, 
            models.StudyHubMember.hub_id == models.StudyHubDeck.hub_id
        ).filter(
            models.StudyHubDeck.deck_id == deck_id, 
            models.StudyHubMember.user_id == user_id
        ).first()
        if not shared:
            raise HTTPException(status_code=403, detail="This deck is private.")
    return _deck_payload(deck)


@app.patch("/api/decks/{deck_id}/visibility")
def update_deck_visibility(deck_id: int, payload: schemas.DeckVisibilityUpdate, db: Session = Depends(get_db)):
    deck = db.query(models.Deck).filter(models.Deck.id == deck_id).first()
    if not deck:
        raise HTTPException(status_code=404, detail="Deck not found.")
    if deck.user_id != payload.user_id:
        raise HTTPException(status_code=403, detail="Only the deck creator can change its visibility.")
    
    vis = "Public" if payload.visibility.lower() == "public" else "Private"
    deck.visibility = vis
    db.commit()
    db.refresh(deck)
    return _deck_payload(deck)


@app.post("/api/decks/{deck_id}/clone")
def clone_deck_to_library(deck_id: int, payload: schemas.DeckCloneRequest, db: Session = Depends(get_db)):
    source_deck = db.query(models.Deck).filter(models.Deck.id == deck_id).first()
    if not source_deck:
        raise HTTPException(status_code=404, detail="Deck not found.")
    if source_deck.visibility != "Public" and source_deck.user_id != payload.user_id:
        raise HTTPException(status_code=403, detail="Cannot clone a private deck.")
    
    cloned_title = f"{source_deck.title} (Saved)" if source_deck.user_id != payload.user_id else f"{source_deck.title} (Copy)"
    cloned_deck = models.Deck(
        title=cloned_title,
        description=source_deck.description,
        visibility="Private",
        user_id=payload.user_id,
        author_name=payload.display_name or "Student",
    )
    db.add(cloned_deck)
    db.commit()
    db.refresh(cloned_deck)

    for card in source_deck.cards:
        cloned_card = models.Card(deck_id=cloned_deck.id, term=card.term, definition=card.definition)
        db.add(cloned_card)
    
    db.commit()
    db.refresh(cloned_deck)
    return _deck_payload(cloned_deck)


@app.post("/api/decks/", response_model=schemas.DeckResponse)
def create_deck(deck: schemas.DeckCreate, db: Session = Depends(get_db)):
    db_deck = models.Deck(
        title=deck.title, 
        description=deck.description, 
        visibility=deck.visibility,
        user_id=deck.user_id,
        author_name=deck.author_name or "Student",
    )
    db.add(db_deck)
    db.commit()
    db.refresh(db_deck)

    for card_data in deck.cards:
        db_card = models.Card(deck_id=db_deck.id, term=card_data.term, definition=card_data.definition)
        db.add(db_card)
    
    db.commit()
    db.refresh(db_deck)
    return db_deck


def _deck_payload(deck):
    return {
        "id": deck.id,
        "title": deck.title,
        "description": deck.description,
        "visibility": deck.visibility or "Private",
        "user_id": deck.user_id,
        "author_name": getattr(deck, "author_name", "Student") or "Student",
        "created_at": deck.created_at,
        "cards": [
            {"id": card.id, "deck_id": card.deck_id, "term": card.term, "definition": card.definition}
            for card in deck.cards
        ],
    }


def _hub_payload(hub, user_id):
    membership = next((member for member in hub.members if member.user_id == user_id), None)
    return {
        "id": hub.id,
        "name": hub.name,
        "description": hub.description,
        "inviteCode": hub.invite_code,
        "ownerId": hub.owner_id,
        "role": membership.role if membership else "Member",
        "members": len(hub.members),
        "decks": len(hub.shared_decks),
        "allowMemberUploads": hub.allow_member_uploads,
        "memberList": [
            {
                "id": member.id,
                "userId": member.user_id,
                "displayName": member.display_name,
                "role": member.role,
                "joinedAt": member.created_at,
            }
            for member in hub.members
        ],
        "sharedDecks": [_deck_payload(shared_deck.deck) for shared_deck in hub.shared_decks],
        "createdAt": hub.created_at,
    }


def _require_hub_member(hub, user_id):
    membership = next((member for member in hub.members if member.user_id == user_id), None)
    if not membership:
        raise HTTPException(status_code=403, detail="Join this hub before accessing it.")
    return membership


def _require_hub_admin(hub, user_id):
    membership = _require_hub_member(hub, user_id)
    if membership.role != "Admin":
        raise HTTPException(status_code=403, detail="Only hub admins can manage members.")
    return membership


@app.get("/api/study-hubs")
def get_study_hubs(user_id: str, db: Session = Depends(get_db)):
    hubs = (
        db.query(models.StudyHub)
        .join(models.StudyHubMember)
        .filter(models.StudyHubMember.user_id == user_id)
        .order_by(models.StudyHub.created_at.desc())
        .all()
    )
    return [_hub_payload(hub, user_id) for hub in hubs]


@app.post("/api/study-hubs")
def create_study_hub(request: schemas.StudyHubCreate, db: Session = Depends(get_db)):
    invite_code = secrets.token_hex(4).upper()
    while db.query(models.StudyHub).filter(models.StudyHub.invite_code == invite_code).first():
        invite_code = secrets.token_hex(4).upper()

    hub = models.StudyHub(
        name=request.name.strip(),
        description=(request.description or "No description provided.").strip(),
        invite_code=invite_code,
        owner_id=request.owner_id,
        owner_name=request.owner_name.strip() or "Student",
        allow_member_uploads=request.allow_member_uploads,
    )
    db.add(hub)
    db.flush()
    db.add(models.StudyHubMember(hub_id=hub.id, user_id=request.owner_id, display_name=hub.owner_name, role="Admin"))

    if request.initial_deck_id:
        deck = db.query(models.Deck).filter(models.Deck.id == request.initial_deck_id).first()
        if not deck:
            raise HTTPException(status_code=404, detail="Initial deck not found.")
        db.add(models.StudyHubDeck(hub_id=hub.id, deck_id=deck.id, shared_by=request.owner_id))

    db.commit()
    db.refresh(hub)
    return _hub_payload(hub, request.owner_id)


@app.get("/api/study-hubs/{hub_id}")
def get_study_hub(hub_id: int, user_id: str, db: Session = Depends(get_db)):
    hub = db.query(models.StudyHub).filter(models.StudyHub.id == hub_id).first()
    if not hub:
        raise HTTPException(status_code=404, detail="Study hub not found.")
    _require_hub_member(hub, user_id)
    return _hub_payload(hub, user_id)


@app.post("/api/study-hubs/join")
def join_study_hub(request: schemas.StudyHubJoin, db: Session = Depends(get_db)):
    hub = db.query(models.StudyHub).filter(models.StudyHub.invite_code == request.invite_code.strip().upper()).first()
    if not hub:
        raise HTTPException(status_code=404, detail="Study hub not found.")

    membership = next((member for member in hub.members if member.user_id == request.user_id), None)
    if not membership:
        db.add(models.StudyHubMember(hub_id=hub.id, user_id=request.user_id, display_name=request.display_name.strip() or "Student", role="Member"))
        db.commit()
        db.refresh(hub)
    return _hub_payload(hub, request.user_id)


@app.post("/api/study-hubs/{hub_id}/decks")
def share_study_hub_deck(hub_id: int, request: schemas.StudyHubDeckShare, db: Session = Depends(get_db)):
    hub = db.query(models.StudyHub).filter(models.StudyHub.id == hub_id).first()
    if not hub:
        raise HTTPException(status_code=404, detail="Study hub not found.")
    membership = _require_hub_member(hub, request.user_id)
    if membership.role != "Admin" and not hub.allow_member_uploads:
        raise HTTPException(status_code=403, detail="Only the hub admin can share decks.")

    deck = db.query(models.Deck).filter(models.Deck.id == request.deck_id).first()
    if not deck:
        raise HTTPException(status_code=404, detail="Deck not found.")
    already_shared = next((shared for shared in hub.shared_decks if shared.deck_id == deck.id), None)
    if not already_shared:
        db.add(models.StudyHubDeck(hub_id=hub.id, deck_id=deck.id, shared_by=request.user_id))
        db.commit()
        db.refresh(hub)
    return _hub_payload(hub, request.user_id)


@app.patch("/api/study-hubs/{hub_id}/settings")
def update_study_hub_settings(hub_id: int, request: schemas.StudyHubPermissionUpdate, db: Session = Depends(get_db)):
    hub = db.query(models.StudyHub).filter(models.StudyHub.id == hub_id).first()
    if not hub:
        raise HTTPException(status_code=404, detail="Study hub not found.")
    _require_hub_admin(hub, request.user_id)
    hub.allow_member_uploads = request.allow_member_uploads
    db.commit()
    db.refresh(hub)
    return _hub_payload(hub, request.user_id)


@app.get("/api/study-hubs/{hub_id}/members")
def get_study_hub_members(hub_id: int, user_id: str, db: Session = Depends(get_db)):
    hub = db.query(models.StudyHub).filter(models.StudyHub.id == hub_id).first()
    if not hub:
        raise HTTPException(status_code=404, detail="Study hub not found.")
    _require_hub_member(hub, user_id)
    return _hub_payload(hub, user_id)["memberList"]


@app.post("/api/study-hubs/{hub_id}/members/{member_id}/remove")
def remove_study_hub_member(hub_id: int, member_id: int, request: schemas.StudyHubMemberAction, db: Session = Depends(get_db)):
    hub = db.query(models.StudyHub).filter(models.StudyHub.id == hub_id).first()
    if not hub:
        raise HTTPException(status_code=404, detail="Study hub not found.")
    _require_hub_admin(hub, request.actor_id)
    member = next((item for item in hub.members if item.id == member_id), None)
    if not member:
        raise HTTPException(status_code=404, detail="Member not found.")
    if member.user_id == hub.owner_id:
        raise HTTPException(status_code=400, detail="The hub owner cannot be removed.")
    db.delete(member)
    db.commit()
    db.refresh(hub)
    return _hub_payload(hub, request.actor_id)


@app.patch("/api/study-hubs/{hub_id}/members/{member_id}/role")
def update_study_hub_member_role(hub_id: int, member_id: int, request: schemas.StudyHubMemberRoleUpdate, db: Session = Depends(get_db)):
    hub = db.query(models.StudyHub).filter(models.StudyHub.id == hub_id).first()
    if not hub:
        raise HTTPException(status_code=404, detail="Study hub not found.")
    if request.actor_id != hub.owner_id:
        raise HTTPException(status_code=403, detail="Only the hub owner can change member roles.")
    if request.role not in {"Admin", "Member"}:
        raise HTTPException(status_code=400, detail="Role must be Admin or Member.")
    member = next((item for item in hub.members if item.id == member_id), None)
    if not member:
        raise HTTPException(status_code=404, detail="Member not found.")
    if member.user_id == hub.owner_id:
        raise HTTPException(status_code=400, detail="The hub owner role cannot be changed.")
    member.role = request.role
    db.commit()
    db.refresh(hub)
    return _hub_payload(hub, request.actor_id)


# --- HYBRID GENERATOR: GROQ (PRIMARY) -> GEMINI (FALLBACK) ---
@app.post("/api/generate-deck/", response_model=schemas.DeckCreate)
def generate_deck_with_ai(request: schemas.GenerateRequest, req: Request, db: Session = Depends(get_db)):
    identifier = get_client_identifier(req, request.user_id)
    enforce_rate_limit("deck_generation", identifier, db)

    # 1. Check in-memory cache for instant response (avoids external API calls & token exhaustion)
    cache_key = make_deck_cache_key(request.topic, request.card_count)
    cached_cards = ai_cache.get(cache_key)
    if cached_cards:
        print(f"Cache HIT for topic: '{request.topic}' ({request.card_count} cards)")
        record_rate_limit_usage("deck_generation", identifier, db)
        return {
            "title": request.topic.title(),
            "description": f"Auto-generated deck covering {request.topic}.",
            "visibility": "Private",
            "user_id": request.user_id,
            "cards": cached_cards
        }

    prompt = f"""
    You are an expert tutor. Generate a highly effective flashcard deck about: {request.topic}.
    Provide exactly {request.card_count} key terms and definitions.
    Respond ONLY with a valid JSON array of objects. Do not include markdown formatting blocks.
    Format exactly like this:
    [
        {{"term": "Concept", "definition": "Explanation"}}
    ]
    """
    
    raw_text = None

    # Concurrency guard: limit simultaneous heavy AI calls to prevent 429 RPM crashes from external AI providers
    with AI_CONCURRENCY_SEMAPHORE:
        # 1. Try Groq First with quick retry
        if groq_client:
            for attempt in range(2):
                try:
                    print(f"Attempting generation with Groq ({GROQ_MODEL})...")
                    chat_completion = groq_client.chat.completions.create(
                        messages=[{"role": "user", "content": prompt}],
                        model=GROQ_MODEL,
                        temperature=0.3,
                    )
                    raw_text = chat_completion.choices[0].message.content
                    print("Groq generation successful!")
                    break
                except Exception as e:
                    print(f"Groq attempt {attempt+1} failed: {e}. Retrying / Falling back...")
                    time.sleep(0.5)

        # 2. Fallback to Gemini if Groq fails or is not configured
        if not raw_text and ai_client:
            for model_name in FALLBACK_GEMINI_MODELS:
                try:
                    print(f"Attempting generation with {model_name}...")
                    response = ai_client.models.generate_content(
                        model=model_name,
                        contents=prompt,
                        config=types.GenerateContentConfig(temperature=0.3)
                    )
                    raw_text = response.text
                    print(f"Success using {model_name}!")
                    break
                except Exception as e:
                    print(f"Model {model_name} overloaded. Falling back...")
                    continue

    if not raw_text:
         raise HTTPException(status_code=503, detail="All AI engines (Groq and Gemini) are currently unavailable.")
         
    try:
        flashcards_data = extract_json_data(raw_text)
        if not isinstance(flashcards_data, list):
            raise ValueError("Expected list of cards")
        
        # Save to cache for future requests across all users
        ai_cache.set(cache_key, flashcards_data)
        record_rate_limit_usage("deck_generation", identifier, db)

        return {
            "title": request.topic.title(),
            "description": f"Auto-generated deck covering {request.topic}.",
            "visibility": "Private",
            "user_id": request.user_id,
            "cards": flashcards_data
        }
        
    except Exception as e:
         raise HTTPException(status_code=500, detail=f"AI returned invalid JSON format: {e}")


def extract_text_from_file_path(file_path: str, ext: str) -> str:
    """Extracts raw text from .pdf, .pptx, or .txt files for fallback LLM processing."""
    ext = ext.lower()
    text_content = ""
    if ext == ".txt":
        for enc in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
            try:
                with open(file_path, "r", encoding=enc, errors="ignore") as f:
                    text_content = f.read()
                if text_content.strip():
                    break
            except Exception:
                continue
    elif ext == ".pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(file_path)
            pages_text = []
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted and extracted.strip():
                    pages_text.append(extracted.strip())
            text_content = "\n\n".join(pages_text)
        except Exception as e:
            print(f"Error extracting text from PDF with pypdf: {e}")
    elif ext == ".pptx":
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            slides_text = []
            for idx, slide in enumerate(prs.slides, 1):
                slide_lines = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            t = paragraph.text.strip()
                            if t:
                                slide_lines.append(t)
                    elif hasattr(shape, "text") and shape.text:
                        t = shape.text.strip()
                        if t:
                            slide_lines.append(t)
                if slide_lines:
                    slides_text.append(f"--- Slide {idx} ---\n" + "\n".join(slide_lines))
            text_content = "\n\n".join(slides_text)
        except Exception as e:
            print(f"Error extracting text from PPTX with python-pptx: {e}")
    return text_content.strip()


def generate_deck_from_text_groq(text_content: str, card_count: int) -> list:
    """Calls Groq to generate flashcards from extracted document text with model fallback and retries."""
    if not groq_client:
        raise ValueError("Groq client is not configured or unavailable.")

    # Truncate text content safely to ~25,000 characters to stay within context and latency limits
    max_chars = 25000
    truncated_text = text_content[:max_chars] if len(text_content) > max_chars else text_content

    prompt = f"""
    You are an expert tutor. Extract the most important concepts from this document content and generate a highly effective flashcard deck.
    Provide exactly {card_count} key terms and definitions based strictly on the document text.
    Respond ONLY with a valid JSON array of objects. Do not include markdown formatting blocks or any explanatory prose.
    Format exactly like this:
    [
        {{"term": "Concept", "definition": "Explanation"}}
    ]

    DOCUMENT CONTENT:
    {truncated_text}
    """

    models_to_try = [GROQ_MODEL]
    if "openai/gpt-oss-120b" in GROQ_MODEL:
        models_to_try.append("openai/gpt-oss-20b")
    elif "openai/gpt-oss-20b" not in models_to_try:
        models_to_try.append("openai/gpt-oss-20b")

    last_err = None
    for model_name in models_to_try:
        for attempt in range(2):
            try:
                print(f"Attempting file extraction fallback with Groq ({model_name}, attempt {attempt + 1})...")
                chat_completion = groq_client.chat.completions.create(
                    messages=[{"role": "user", "content": prompt}],
                    model=model_name,
                    temperature=0.3,
                )
                raw_text = chat_completion.choices[0].message.content
                print(f"Groq fallback file generation successful using {model_name}!")
                data = extract_json_data(raw_text)
                if isinstance(data, list) and len(data) > 0:
                    return data
                raise ValueError("AI response did not contain a valid non-empty list of cards")
            except Exception as e:
                last_err = e
                print(f"Groq ({model_name}) attempt {attempt + 1} failed: {e}")
                time.sleep(0.5)

    raise ValueError(f"Groq generation failed: {last_err}")


# --- FILE GENERATOR: GEMINI WITH AUTOMATIC GROQ FALLBACK ---
@app.post("/api/generate-deck-from-file/", response_model=schemas.DeckCreate)
async def generate_deck_from_file(req: Request, file: UploadFile = File(...), card_count: int = Form(10), user_id: str | None = Form(None), db: Session = Depends(get_db)):
    identifier = get_client_identifier(req, user_id)
    enforce_rate_limit("deck_generation", identifier, db)

    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in ['.pdf', '.txt', '.pptx']:
        raise HTTPException(status_code=400, detail="Unsupported file format. Please upload a .pdf, .pptx, or .txt file.")
        
    with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as temp_file:
        content = await file.read()
        temp_file.write(content)
        temp_file_path = temp_file.name

    flashcards_data = None
    gemini_error = None

    try:
        # 1. Primary engine: Attempt Gemini native document processing
        if ai_client:
            uploaded_gemini_file = None
            try:
                uploaded_gemini_file = ai_client.files.upload(file=temp_file_path)
                
                prompt = f"""
                You are an expert tutor. Extract the most important concepts from this uploaded document and generate a highly effective flashcard deck.
                Provide exactly {card_count} key terms and definitions based strictly on the document content.
                Respond ONLY with a valid JSON array of objects. Do not include markdown formatting blocks.
                Format exactly like this:
                [
                    {{"term": "Concept", "definition": "Explanation"}}
                ]
                """
                
                response = None
                with AI_CONCURRENCY_SEMAPHORE:
                    for model_name in FALLBACK_GEMINI_MODELS:
                        try:
                            print(f"Attempting file extraction with {model_name}...")
                            response = ai_client.models.generate_content(
                                model=model_name,
                                contents=[uploaded_gemini_file, prompt],
                                config=types.GenerateContentConfig(temperature=0.3)
                            )
                            print(f"Success using {model_name}!")
                            break
                        except Exception as e:
                            print(f"Model {model_name} overloaded. Falling back... ({e})")
                            continue

                if response and response.text:
                    parsed = extract_json_data(response.text)
                    if isinstance(parsed, list) and len(parsed) > 0:
                        flashcards_data = parsed

            except Exception as e:
                gemini_error = e
                print(f"Gemini file processing error: {e}")
            finally:
                if uploaded_gemini_file:
                    try:
                        ai_client.files.delete(name=uploaded_gemini_file.name)
                    except Exception:
                        pass
        else:
            gemini_error = "Gemini client not initialized"

        # 2. Secondary engine: If Gemini failed or is unavailable, switch to Groq fallback
        if not flashcards_data:
            print(f"Gemini file generation failed ({gemini_error}). Switching to Groq fallback...")
            if not groq_client:
                raise HTTPException(
                    status_code=503,
                    detail=f"Gemini failed ({gemini_error}) and Groq client is not available."
                )

            extracted_text = extract_text_from_file_path(temp_file_path, ext)
            if not extracted_text:
                raise HTTPException(
                    status_code=400,
                    detail="Could not extract readable text from document. Ensure the file contains selectable text and is not an empty or scanned document."
                )

            with AI_CONCURRENCY_SEMAPHORE:
                try:
                    flashcards_data = generate_deck_from_text_groq(extracted_text, card_count)
                except Exception as e:
                    print(f"Groq fallback file generation failed: {e}")
                    raise HTTPException(
                        status_code=503,
                        detail=f"Both Gemini ({gemini_error}) and Groq ({e}) failed to generate the deck."
                    )

        if not flashcards_data or not isinstance(flashcards_data, list):
            raise HTTPException(status_code=500, detail="AI returned invalid card format.")

        clean_title = file.filename.replace(ext, "").replace("-", " ").replace("_", " ").title()
        record_rate_limit_usage("deck_generation", identifier, db)

        return {
            "title": clean_title,
            "description": f"Extracted from document: {file.filename}",
            "visibility": "Private",
            "user_id": user_id,
            "cards": flashcards_data
        }

    finally:
        if os.path.exists(temp_file_path):
            try:
                os.remove(temp_file_path)
            except Exception:
                pass

@app.post("/api/decks/{deck_id}/readout")
def generate_deck_readout(deck_id: int, req: Request, user_id: str | None = None, db: Session = Depends(get_db)):
    identifier = get_client_identifier(req, user_id)
    enforce_rate_limit("readout", identifier, db)

    deck = db.query(models.Deck).filter(models.Deck.id == deck_id).first()
    if not deck:
        raise HTTPException(status_code=404, detail="Deck not found.")
    if not deck.cards:
        raise HTTPException(status_code=400, detail="Add cards to this deck before generating a readout.")
    if not groq_client and not ai_client:
        raise HTTPException(status_code=503, detail="AI services are currently unavailable.")

    card_context = "\n".join(
        f"- {card.term}: {card.definition}" for card in deck.cards[:200]
    )
    prompt = f"""
    You are Zero, an expert study-material author. Create a clear reading guide from the flashcards below.
    The guide must teach the material, connect related concepts, and help a student prepare for an assessment.
    Do not invent facts that are not supported by the cards.

    Return ONLY valid JSON with this exact shape:
    {{
      "title": "A concise title",
      "overview": "A two or three sentence overview",
      "sections": [
        {{"heading": "Concept group", "body": "A concise explanation with examples where supported by the cards."}}
      ],
      "key_takeaways": ["Short takeaway", "Short takeaway"]
    }}
    Create 3 to 6 sections and 4 to 8 key takeaways.

    Deck title: {deck.title}
    Deck description: {deck.description or "No description provided."}
    Flashcards:
    {card_context}
    """

    readout_cache_key = make_readout_cache_key(deck.id, len(deck.cards), deck.title)
    cached_readout = ai_cache.get(readout_cache_key)
    if cached_readout:
        readout = cached_readout
    else:
        raw_text = None
        with AI_CONCURRENCY_SEMAPHORE:
            if groq_client:
                try:
                    response = groq_client.chat.completions.create(
                        messages=[{"role": "user", "content": prompt}],
                        model=GROQ_MODEL,
                        temperature=0.3,
                    )
                    raw_text = response.choices[0].message.content
                except Exception as error:
                    print(f"Groq readout generation failed: {error}. Falling back to Gemini...")

            if not raw_text and ai_client:
                for model_name in FALLBACK_GEMINI_MODELS:
                    try:
                        response = ai_client.models.generate_content(
                            model=model_name,
                            contents=prompt,
                            config=types.GenerateContentConfig(temperature=0.3),
                        )
                        raw_text = response.text
                        break
                    except Exception as error:
                        print(f"{model_name} readout generation failed: {error}")

        if not raw_text:
            raise HTTPException(status_code=503, detail="AI services are currently unavailable.")

        try:
            readout = extract_json_data(raw_text)
            ai_cache.set(readout_cache_key, readout, ttl=86400)
        except Exception:
            raise HTTPException(status_code=502, detail="AI returned an invalid reading guide.")

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReadoutTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=24,
        leading=30,
        textColor=HexColor("#211C18"),
        alignment=TA_CENTER,
        spaceAfter=12,
    )
    subtitle_style = ParagraphStyle(
        "ReadoutSubtitle",
        parent=styles["Normal"],
        fontSize=10,
        leading=14,
        textColor=HexColor("#997E67"),
        alignment=TA_CENTER,
        spaceAfter=24,
    )
    heading_style = ParagraphStyle(
        "ReadoutHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=19,
        textColor=HexColor("#664930"),
        spaceBefore=16,
        spaceAfter=7,
    )
    body_style = ParagraphStyle(
        "ReadoutBody",
        parent=styles["BodyText"],
        fontSize=10.5,
        leading=16,
        textColor=HexColor("#3F342D"),
        spaceAfter=9,
    )
    takeaway_style = ParagraphStyle(
        "ReadoutTakeaway",
        parent=body_style,
        leftIndent=14,
        firstLineIndent=-10,
    )

    def safe_paragraph(value):
        from xml.sax.saxutils import escape
        replacements = str.maketrans({
            "\u2010": "-",
            "\u2011": "-",
            "\u2012": "-",
            "\u2013": "-",
            "\u2014": "-",
            "\u2018": "'",
            "\u2019": "'",
            "\u201c": '"',
            "\u201d": '"',
            "\u2026": "...",
            "\u2022": "-",
            "\u00a0": " ",
        })
        normalized = str(value or "").translate(replacements)
        normalized = normalized.encode("latin-1", "replace").decode("latin-1")
        return escape(normalized).replace("\n", "<br/>")

    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=0.7 * inch,
        leftMargin=0.7 * inch,
        topMargin=0.65 * inch,
        bottomMargin=0.65 * inch,
        title=f"{deck.title} Reading Guide",
        author="CramZero Zero AI",
    )
    story = [
        Paragraph(safe_paragraph(readout.get("title") or f"{deck.title} Reading Guide"), title_style),
        Paragraph("Generated by Zero AI from your CramZero deck", subtitle_style),
        Paragraph("Overview", heading_style),
        Paragraph(safe_paragraph(readout.get("overview")), body_style),
    ]
    for section in readout.get("sections", []):
        story.extend([
            Paragraph(safe_paragraph(section.get("heading")), heading_style),
            Paragraph(safe_paragraph(section.get("body")), body_style),
        ])
    story.append(Paragraph("Key Takeaways", heading_style))
    for takeaway in readout.get("key_takeaways", []):
        story.append(Paragraph(f"- {safe_paragraph(takeaway)}", takeaway_style))

    document.build(story, onFirstPage=lambda canvas, doc: None, onLaterPages=lambda canvas, doc: None)
    buffer.seek(0)
    filename = re.sub(r"[^A-Za-z0-9]+", "-", deck.title).strip("-").lower() or "cramzero-readout"
    record_rate_limit_usage("readout", identifier, db)
    return StreamingResponse(
        iter([buffer.getvalue()]),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}-reading-guide.pdf"'},
    )

@app.post("/api/chat/")
def ai_tutor_chat(request: schemas.ChatRequest, req: Request, db: Session = Depends(get_db)):
    identifier = get_client_identifier(req, request.user_id)
    enforce_rate_limit("chat", identifier, db)

    cache_key = make_chat_cache_key(request.context, request.message)
    cached_reply = ai_cache.get(cache_key)
    if cached_reply:
        record_rate_limit_usage("chat", identifier, db)
        return {"reply": cached_reply}

    prompt = f"""
    You are Zero, an expert, encouraging, and highly organized AI tutor for the CramZero study app.
    You have full access to the user's current flashcard deck context below.
    
    CRITICAL FORMATTING RULES:
    1. ALWAYS structure your responses using Markdown.
    2. Break down explanations into scannable bullet points or numbered lists.
    3. Use bold text to emphasize key terms.
    4. Use ### headers to separate different ideas or steps.
    5. Never send a giant wall of plain text. Keep it visually digestible.

    --- DECK CONTEXT ---
    {request.context}
    --- END CONTEXT ---

    User's message: {request.message}
    """
    reply_text = None
    
    with AI_CONCURRENCY_SEMAPHORE:
        if groq_client:
            try:
                chat_completion = groq_client.chat.completions.create(
                    messages=[{"role": "user", "content": prompt}],
                    model=GROQ_MODEL,
                    temperature=0.4,
                )
                reply_text = chat_completion.choices[0].message.content
            except Exception as e:
                print(f"Groq chat failed: {e}. Falling back to Gemini...")
                
        # 2. Fallback to Gemini Matrix
        if not reply_text and ai_client:
            for model_name in FALLBACK_GEMINI_MODELS:
                try:
                    response = ai_client.models.generate_content(
                        model=model_name,
                        contents=prompt,
                        config=types.GenerateContentConfig(temperature=0.4)
                    )
                    reply_text = response.text
                    break
                except Exception:
                    continue
                
    if not reply_text:
        raise HTTPException(status_code=503, detail="Professor AI is currently unavailable.")
        
    ai_cache.set(cache_key, reply_text, ttl=43200)
    record_rate_limit_usage("chat", identifier, db)
    return {"reply": reply_text}


# ============================================================================
# MULTIPLAYER REAL-TIME BATTLE WEBSOCKET HANDLERS
# ============================================================================

QUESTION_SECONDS = 15
multiplayer_rooms = {}  # pin -> room dict
_multiplayer_ticker_started = False


def make_pin() -> str:
    while True:
        pin = str(secrets.randbelow(900000) + 100000)
        if pin not in multiplayer_rooms:
            return pin


def make_player_id() -> str:
    return str(uuid.uuid4())


def public_room_state(room: dict) -> dict:
    phase = room["phase"]
    deck = dict(room.get("deck") or {})
    cards = deck.get("cards", [])
    safe_cards = []
    include_def = phase in ("round-result", "results", "cancelled")
    for card in cards:
        card_copy = dict(card)
        if not include_def:
            card_copy.pop("definition", None)
        safe_cards.append(card_copy)
    deck["cards"] = safe_cards

    now = int(time.time() * 1000)
    countdown_rem = 0
    if phase == "countdown" and room.get("countdownStartedAt"):
        countdown_rem = max(1, int(math.ceil((3000 - (now - room["countdownStartedAt"])) / 1000.0)))
    question_rem = QUESTION_SECONDS
    if phase == "playing" and room.get("questionStartedAt"):
        question_rem = max(0, int(math.ceil((QUESTION_SECONDS * 1000 - (now - room["questionStartedAt"])) / 1000.0)))

    players_safe = [
        {
            "id": p["id"],
            "name": p["name"],
            "avatar": p.get("avatar") or "",
            "score": p["score"],
            "isHost": p.get("isHost", False),
            "stats": p.get("stats") or {"correct": 0, "total": 0, "streak": 0, "maxStreak": 0},
        }
        for p in room["players"]
    ]

    return {
        "type": "room_state",
        "room": {
            "pin": room["pin"],
            "phase": room["phase"],
            "hostId": room["hostId"],
            "deck": deck,
            "players": players_safe,
            "questionIndex": room["questionIndex"],
            "questionStartedAt": room["questionStartedAt"],
            "countdownStartedAt": room["countdownStartedAt"],
            "serverTime": now,
            "countdownRemaining": countdown_rem,
            "questionSecondsRemaining": question_rem,
            "cancelReason": room.get("cancelReason"),
            "questionSeconds": QUESTION_SECONDS,
            "answeredPlayerIds": list(room["answers"].keys()),
            "scores": {p["id"]: p["score"] for p in room["players"]},
        },
    }


async def send_ws(ws: WebSocket, message: dict):
    try:
        await ws.send_json(message)
    except Exception:
        pass


async def broadcast_room(room: dict, message: dict):
    for p in list(room["players"]):
        ws = p.get("ws")
        if ws:
            await send_ws(ws, message)


async def finish_round(room: dict):
    if room["phase"] != "playing":
        return
    room["phase"] = "round-result"
    room["questionStartedAt"] = None
    room["_lastPlayingTick"] = None
    room["_lastCountdownTick"] = None
    await broadcast_room(room, public_room_state(room))


async def start_question(room: dict):
    room["phase"] = "playing"
    room["answers"] = {}
    room["questionStartedAt"] = int(time.time() * 1000)
    room["countdownStartedAt"] = None
    room["_lastPlayingTick"] = QUESTION_SECONDS
    room["_lastCountdownTick"] = None
    await broadcast_room(room, public_room_state(room))
    await broadcast_room(room, {
        "type": "timer_tick",
        "phase": "playing",
        "timeLeft": QUESTION_SECONDS,
        "serverTime": room["questionStartedAt"]
    })


async def remove_player(room: dict, player_id: str):
    was_active_match = room["phase"] in ("countdown", "playing", "round-result")
    room["players"] = [p for p in room["players"] if p["id"] != player_id]
    room["answers"].pop(player_id, None)

    if len(room["players"]) == 0:
        multiplayer_rooms.pop(room["pin"], None)
        return

    if was_active_match and len(room["players"]) < 2:
        room["phase"] = "cancelled"
        room["questionStartedAt"] = None
        room["countdownStartedAt"] = None
        room["cancelReason"] = "The battle ended because the other player left."
        await broadcast_room(room, public_room_state(room))
        return

    if room["hostId"] == player_id:
        room["hostId"] = room["players"][0]["id"]
        room["players"][0]["isHost"] = True

    await broadcast_room(room, public_room_state(room))


async def multiplayer_tick_loop():
    while True:
        try:
            await asyncio.sleep(0.2)
            now = int(time.time() * 1000)
            for room in list(multiplayer_rooms.values()):
                phase = room.get("phase")
                if phase == "countdown" and room.get("countdownStartedAt"):
                    elapsed = now - room["countdownStartedAt"]
                    if elapsed >= 3000:
                        await start_question(room)
                    else:
                        remaining = max(1, int(math.ceil((3000 - elapsed) / 1000.0)))
                        if room.get("_lastCountdownTick") != remaining:
                            room["_lastCountdownTick"] = remaining
                            await broadcast_room(room, {
                                "type": "timer_tick",
                                "phase": "countdown",
                                "countdown": remaining,
                                "serverTime": now
                            })
                elif phase == "playing" and room.get("questionStartedAt"):
                    elapsed = now - room["questionStartedAt"]
                    if elapsed >= QUESTION_SECONDS * 1000:
                        await finish_round(room)
                    else:
                        remaining = max(0, int(math.ceil((QUESTION_SECONDS * 1000 - elapsed) / 1000.0)))
                        if room.get("_lastPlayingTick") != remaining:
                            room["_lastPlayingTick"] = remaining
                            await broadcast_room(room, {
                                "type": "timer_tick",
                                "phase": "playing",
                                "timeLeft": remaining,
                                "serverTime": now
                            })
        except asyncio.CancelledError:
            break
        except Exception:
            pass


def ensure_ticker_running():
    global _multiplayer_ticker_started
    if not _multiplayer_ticker_started:
        try:
            loop = asyncio.get_running_loop()
            loop.create_task(multiplayer_tick_loop())
            _multiplayer_ticker_started = True
        except RuntimeError:
            pass


@app.on_event("startup")
async def startup_multiplayer_ticker():
    ensure_ticker_running()


@app.websocket("/ws/multiplayer")
async def websocket_multiplayer(websocket: WebSocket):
    await websocket.accept()
    ensure_ticker_running()
    current_room_pin = None
    current_player_id = None
    try:
        while True:
            text_data = await websocket.receive_text()
            try:
                payload = json.loads(text_data)
            except Exception:
                await send_ws(websocket, {"type": "error", "message": "Invalid multiplayer message."})
                continue

            msg_type = payload.get("type")

            if msg_type == "host":
                name = str(payload.get("name") or "").strip()[:32]
                avatar = str(payload.get("avatar") or "").strip()
                deck = payload.get("deck") or {}
                cards = deck.get("cards") or []
                if not name or not cards:
                    await send_ws(websocket, {"type": "error", "message": "A player name and playable deck are required."})
                    continue

                player_id = make_player_id()
                pin = make_pin()
                room = {
                    "pin": pin,
                    "phase": "lobby",
                    "hostId": player_id,
                    "deck": deck,
                    "players": [
                        {
                            "id": player_id,
                            "name": name,
                            "avatar": avatar,
                            "score": 0,
                            "isHost": True,
                            "stats": {"correct": 0, "total": 0, "streak": 0, "maxStreak": 0},
                            "ws": websocket,
                        }
                    ],
                    "questionIndex": 0,
                    "questionStartedAt": None,
                    "countdownStartedAt": None,
                    "cancelReason": None,
                    "answers": {},
                }
                multiplayer_rooms[pin] = room
                current_room_pin = pin
                current_player_id = player_id

                resp = {**public_room_state(room), "type": "connected", "playerId": player_id}
                await send_ws(websocket, resp)
                continue

            if msg_type == "join":
                pin = re.sub(r"\D", "", str(payload.get("pin") or ""))[:6]
                room = multiplayer_rooms.get(pin)
                if not room:
                    await send_ws(websocket, {"type": "error", "message": "That lobby was not found."})
                    continue
                if room["phase"] != "lobby":
                    await send_ws(websocket, {"type": "error", "message": "That battle has already started."})
                    continue
                name = str(payload.get("name") or "").strip()[:32]
                avatar = str(payload.get("avatar") or "").strip()
                if not name:
                    await send_ws(websocket, {"type": "error", "message": "Enter a display name to join."})
                    continue

                player_id = make_player_id()
                room["players"].append({
                    "id": player_id,
                    "name": name,
                    "avatar": avatar,
                    "score": 0,
                    "isHost": False,
                    "stats": {"correct": 0, "total": 0, "streak": 0, "maxStreak": 0},
                    "ws": websocket,
                })
                current_room_pin = pin
                current_player_id = player_id

                resp = {**public_room_state(room), "type": "connected", "playerId": player_id}
                await send_ws(websocket, resp)
                await broadcast_room(room, public_room_state(room))
                continue

            if not current_room_pin or not current_player_id:
                await send_ws(websocket, {"type": "error", "message": "Join or host a lobby first."})
                continue

            room = multiplayer_rooms.get(current_room_pin)
            if not room:
                await send_ws(websocket, {"type": "error", "message": "Lobby no longer exists."})
                continue

            player = next((p for p in room["players"] if p["id"] == current_player_id), None)
            if not player:
                continue

            if msg_type == "leave":
                room_to_leave = room
                pid = current_player_id
                current_room_pin = None
                current_player_id = None
                await remove_player(room_to_leave, pid)
                await websocket.close()
                break

            if msg_type == "start":
                if current_player_id != room["hostId"] or len(room["players"]) < 2 or room["phase"] != "lobby":
                    continue
                room["phase"] = "countdown"
                room["questionIndex"] = 0
                room["answers"] = {}
                room["questionStartedAt"] = None
                room["countdownStartedAt"] = int(time.time() * 1000)
                room["_lastCountdownTick"] = 3
                room["_lastPlayingTick"] = None
                await broadcast_room(room, public_room_state(room))
                await broadcast_room(room, {
                    "type": "timer_tick",
                    "phase": "countdown",
                    "countdown": 3,
                    "serverTime": room["countdownStartedAt"]
                })
                continue

            if msg_type == "answer":
                if room["phase"] != "playing" or current_player_id in room["answers"]:
                    continue
                cards = room["deck"].get("cards", [])
                if room["questionIndex"] >= len(cards):
                    continue
                card = cards[room["questionIndex"]]
                now_ms = int(time.time() * 1000)
                q_start = room.get("questionStartedAt") or now_ms
                elapsed = max(0, now_ms - q_start)
                remaining = max(0, QUESTION_SECONDS * 1000 - elapsed)
                is_correct = (payload.get("answer") == card.get("definition"))
                points = round(1000 * (remaining / (QUESTION_SECONDS * 1000))) if is_correct else 0
                player["score"] += points

                stats = player.setdefault("stats", {"correct": 0, "total": 0, "streak": 0, "maxStreak": 0})
                stats["total"] += 1
                if is_correct:
                    stats["correct"] += 1
                    stats["streak"] += 1
                    stats["maxStreak"] = max(stats["maxStreak"], stats["streak"])
                else:
                    stats["streak"] = 0

                room["answers"][current_player_id] = {
                    "answer": payload.get("answer"),
                    "isCorrect": is_correct,
                    "points": points,
                    "reactionTime": round(elapsed / 1000, 2),
                }
                await send_ws(websocket, {
                    "type": "answer_result",
                    "isCorrect": is_correct,
                    "points": points,
                    "correctAnswer": card.get("definition"),
                })
                await broadcast_room(room, public_room_state(room))
                if len(room["answers"]) == len(room["players"]):
                    await finish_round(room)
                continue

            if msg_type == "next_question":
                if current_player_id != room["hostId"] or room["phase"] != "round-result":
                    continue
                cards = room["deck"].get("cards", [])
                if room["questionIndex"] >= len(cards) - 1:
                    room["phase"] = "results"
                    await broadcast_room(room, public_room_state(room))
                    continue
                room["questionIndex"] += 1
                room["answers"] = {}
                room["phase"] = "countdown"
                room["questionStartedAt"] = None
                room["countdownStartedAt"] = int(time.time() * 1000)
                room["_lastCountdownTick"] = 3
                room["_lastPlayingTick"] = None
                await broadcast_room(room, public_room_state(room))
                await broadcast_room(room, {
                    "type": "timer_tick",
                    "phase": "countdown",
                    "countdown": 3,
                    "serverTime": room["countdownStartedAt"]
                })
                continue

            if msg_type == "rematch":
                if current_player_id != room["hostId"] or room["phase"] != "results":
                    continue
                for p in room["players"]:
                    p["score"] = 0
                    p["stats"] = {"correct": 0, "total": 0, "streak": 0, "maxStreak": 0}
                room["questionIndex"] = 0
                room["answers"] = {}
                room["phase"] = "countdown"
                room["questionStartedAt"] = None
                room["countdownStartedAt"] = int(time.time() * 1000)
                room["_lastCountdownTick"] = 3
                room["_lastPlayingTick"] = None
                await broadcast_room(room, public_room_state(room))
                await broadcast_room(room, {
                    "type": "timer_tick",
                    "phase": "countdown",
                    "countdown": 3,
                    "serverTime": room["countdownStartedAt"]
                })
                continue

    except WebSocketDisconnect:
        pass
    except Exception as e:
        print(f"Multiplayer WebSocket connection error: {e}")
    finally:
        if current_room_pin and current_player_id:
            room = multiplayer_rooms.get(current_room_pin)
            if room:
                await remove_player(room, current_player_id)

